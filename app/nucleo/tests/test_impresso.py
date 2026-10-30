"""S13: tela do impresso ("o que mudou"), "Revista Impressa Atualizada" e planilhas (D-072, D-074, D-089, D-094)."""

import io
import json
import zipfile
from pathlib import Path

import openpyxl
from django.db import transaction
from django.utils import timezone
from django.test import SimpleTestCase
from django.urls import reverse

from nucleo.models import Corte, Descritor, FotoRevista, Projeto, ProjetoPasso, Revista, Tarefa, TemplateVersao
from nucleo.models.revista import Motivo
from nucleo.servicos import bloqueio, planilhas, validacao
from nucleo.servicos.foto import fotografar
from nucleo import texto
from nucleo.servicos.resolver import carregar, hash_validado, resolver
from nucleo.tests.test_admin_projeto import Base, mensagens
from nucleo.tests.test_painel import SVG_B, completar, trocar_svg
from nucleo.tests.test_resolver import F1

FIXTURES = Path(__file__).parent / "fixtures"
CORES_ANTIGAS = {"padrao-1": "#f72628", "padrao-2": "#ff8d43", "padrao-3": "#b3ccff", "padrao-4": "#0ed3ff"}  # CORES_PADRAO do TS


class ImpressoTest(Base):
    def setUp(self):
        super().setUp()
        self.rev = self.go["LP"]
        completar(self.rev)
        self.url = reverse("impresso", args=[self.rev.pk])

    def validar(self, revista=None):
        revista = revista or self.rev
        with transaction.atomic():
            self.assertIsNone(validacao.validar(revista.pk, hash_validado(carregar(revista.pk)), self.ped))

    def abrir(self):
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 200)
        return r

    def marcar(self, hash_visto=None):
        hash_visto = hash_visto or self.abrir().context["hash"]
        return self.client.post(reverse("impresso_marcar", args=[self.rev.pk]), {"hash_visto": hash_visto}, follow=True)

    def grupos(self):
        """{fascículo: [grupo]} do que mudou na tela."""
        return {f: list(g) for f, g in self.abrir().context["mudancas"].items()}

    def imprimir(self):
        self.validar()
        self.client.force_login(self.imp)
        self.assertIn("Revista Impressa Atualizada: os avisos desta revista foram zerados.", mensagens(self.marcar()))

    def test_primeira_entrega_tudo_novo(self):
        self.client.force_login(self.imp)
        r = self.abrir()
        self.assertContains(r, "Primeira entrega")
        self.assertEqual(self.grupos(), {"F1": ["introducao", "Cores", "5EF", "9EF", "fechamento"], "F2": ["abertura", "passo1", "dlg-p2"]})
        itens = [i for g in r.context["mudancas"].values() for lista in g.values() for i in lista]
        self.assertTrue(all(i["antes"] is None for i in itens))
        cinco = [i["item"] for i in r.context["mudancas"]["F1"]["5EF"]]
        self.assertEqual(cinco, ["Escala", "D01", "D02"])  # ordem da página
        escala = r.context["mudancas"]["F1"]["5EF"][0]["depois"]
        self.assertEqual(escala.splitlines()[0], "Padrão 01: 0 a 100")
        self.assertIn("Copiar texto novo", r.content.decode())
        self.assertNotContains(r, 'id="botao-impressa"')  # sem "Validado"
        self.assertContains(r, "Aguardando o “Validado”")

    def test_marca_e_zera(self):
        self.imprimir()
        foto = FotoRevista.objects.get(revista=self.rev)
        self.assertEqual((foto.motivo, foto.criada_por, foto.hash), (Motivo.IMPRESSO, self.imp, hash_validado(carregar(self.rev.pk))))
        r = self.abrir()
        self.assertEqual(r.context["mudancas"], {})
        self.assertContains(r, "Sem mudanças desde a última")
        self.assertNotContains(r, 'id="botao-impressa"')
        painel = self.client.get(reverse("painel", args=[2026]))
        linha = next(x for x in painel.context["linhas"] if x["revista"].pk == self.rev.pk)
        b = f'{reverse("ui_impresso", args=[2026])}?revista={self.rev.pk}'  # a coluna leva à tela B
        self.assertEqual(linha["faltas"]["impresso"], {"ok": True, "dica": "", "link": b})
        self.assertContains(painel, f'<a href="{b}" title="Ver a última versão impressa">')  # ✓ também é link

    def test_variavel_aparece_em_todas_as_secoes_que_a_usam(self):
        rf_template = self.templates["LP", "F1"]
        TemplateVersao.objects.create(template=rf_template, numero=2, texto=F1.replace("Fim.", "Fim de {{nome_projeto}}."))
        self.imprimir()
        self.projeto_go.nome = "Projeto GO novo"
        self.projeto_go.save()
        self.assertEqual(self.grupos(), {"F1": ["introducao", "fechamento"]})
        item = self.abrir().context["mudancas"]["F1"]["fechamento"][0]
        self.assertIn("Fim de Projeto GO.", item["antes"])
        self.assertIn("Fim de Projeto GO novo.", item["depois"])
        self.assertIn('<span class="diff_add">&nbsp;novo</span>', item["tabela"])

    def test_corte_alterado_aparece_so_na_etapa(self):
        self.imprimir()
        Corte.objects.filter(parceiro_ano=self.projeto_go.parceiro_ano, etapa=self.e5, componente=self.lp).update(valores=[100, 200, 260, 300])
        self.assertEqual(self.grupos(), {"F1": ["5EF"]})

    def test_nome_do_padrao_aparece_em_todas_as_etapas(self):  # D-175
        self.imprimir()
        self.projeto_go.pd02 = "Básico"
        self.projeto_go.save()
        mudou = self.abrir().context["mudancas"]["F1"]
        self.assertEqual(list(mudou), ["5EF", "9EF"])
        self.assertEqual([i["item"] for i in mudou["9EF"]], ["Escala"])
        self.assertIn("Básico: 100 a 200", mudou["9EF"][0]["depois"])

    def test_hash_velho_recusado(self):
        self.validar()
        self.client.force_login(self.imp)
        aberta = self.abrir()
        self.assertContains(aberta, 'id="botao-impressa"')
        trocar_svg(self.rev, SVG_B)  # mudou depois que a página abriu e a pedagógica validou a nova versão
        self.validar()
        r = self.marcar(aberta.context["hash"])
        self.assertIn("Impresso não atualizado: a revista não está “Validado” na versão que a página mostrou; recarregue.", mensagens(r))
        self.assertFalse(FotoRevista.objects.exists())

    def test_mudou_depois_do_validado_recusa_e_desfaz_a_foto(self):
        self.validar()
        self.client.force_login(self.imp)
        visto = self.abrir().context["hash"]
        trocar_svg(self.rev, SVG_B)  # hash_validado continua = visto, mas a revista atual é outra
        self.assertIn("Impresso não atualizado: a revista mudou depois que a página abriu; confira de novo.", mensagens(self.marcar(visto)))
        self.assertFalse(FotoRevista.objects.exists())

    def test_projeto_em_edicao_recusa(self):
        self.validar()
        self.client.force_login(self.imp)
        aberta = self.abrir()
        self.travar(self.projeto_go)
        self.assertIn("Impresso não atualizado: projeto em edição por outro.", mensagens(self.marcar(aberta.context["hash"])))
        self.assertFalse(FotoRevista.objects.exists())
        self.assertEqual(bloqueio.quem_bloqueia(self.projeto_go).usuario, self.outro)

    def test_readicionada_compara_com_a_foto_herdada(self):
        # A cascata copia a última foto do arquivado para a revista nova (test_cascata); aqui, a tela.
        foto = fotografar(carregar(self.rev.pk), Motivo.IMPRESSO, self.usuario)
        FotoRevista.objects.filter(pk=foto.pk).update(herdada=True)
        self.projeto_go.adjetivo_patrio_plural_masculino = "goianienses"
        self.projeto_go.save()
        r = self.abrir()
        self.assertContains(r, "foto herdada")
        self.assertEqual({f: list(g) for f, g in r.context["mudancas"].items()}, {"F2": ["passo1"]})

    def test_permissoes(self):
        url = reverse("impresso_marcar", args=[self.rev.pk])
        hash_atual = hash_validado(carregar(self.rev.pk))
        # Pedagógica vê a tela, mas não marca.
        self.validar()
        self.assertNotContains(self.abrir(), 'id="botao-impressa"')
        self.assertEqual(self.client.post(url, {"hash_visto": hash_atual}).status_code, 403)
        # Impresso: só POST.
        self.client.force_login(self.imp)
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertFalse(FotoRevista.objects.exists())
        # Impresso antes do "Validado": recusado.
        mt = self.go["MT"]
        completar(mt)
        r = self.client.post(reverse("impresso_marcar", args=[mt.pk]), {"hash_visto": hash_validado(carregar(mt.pk))}, follow=True)
        self.assertIn("Impresso não atualizado: a revista não está “Validado” na versão que a página mostrou; recarregue.", mensagens(r))
        self.assertFalse(FotoRevista.objects.exists())
        # Anônimo vai ao login.
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)

    def nova_versao_f1(self, texto_f1):
        t = self.templates["LP", "F1"]
        TemplateVersao.objects.create(template=t, numero=TemplateVersao.objects.filter(template=t).count() + 1, texto=texto_f1)

    def test_so_o_fim_da_tarefa_aparece_no_diff(self):
        self.imprimir()
        Tarefa.objects.filter(pk=self.tarefa.pk).update(fim=175)  # mesmo padrão: só o fim muda
        mudou = self.abrir().context["mudancas"]
        self.assertEqual({f: list(g) for f, g in mudou.items()}, {"F1": ["5EF"]})
        item = mudou["F1"]["5EF"][0]
        self.assertEqual(item["item"], "D01")
        self.assertIn("  100–150: Localiza.", item["antes"])
        self.assertIn("  100–175: Localiza.", item["depois"])

    def test_item_da_tarefa_no_diff(self):
        self.imprimir()
        Tarefa.objects.filter(pk=self.tarefa.pk).update(arquivo="itens/abc.pdf")
        self.assertIn("  100–150: Localiza. [item: abc.pdf]", self.abrir().context["mudancas"]["F1"]["5EF"][0]["depois"])

    def test_so_formatacao_aparece_no_diff(self):
        self.imprimir()
        self.nova_versao_f1(F1.replace("**{{nome_projeto}}**", "{{nome_projeto}}"))
        item = self.abrir().context["mudancas"]["F1"]["introducao"][0]
        self.assertIn("Texto **Projeto GO**.", item["antes"])
        self.assertIn("Texto Projeto GO.", item["depois"])

    def test_lista_vira_marcacao(self):
        self.assertEqual(
            texto.html_para_texto("<p>a</p>\n<ol>\n<li>um</li>\n<li><em>dois</em></li>\n</ol>\n<ul>\n<li>x</li>\n</ul>\n\n\n<p>b</p>"),
            "a\n\n1. um\n2. *dois*\n\n- x\n\nb",  # listas separadas: linha em branco entre elas
        )

    def test_secao_removida(self):
        self.imprimir()
        self.nova_versao_f1(F1.split("## fechamento")[0])
        r = self.abrir()
        item = r.context["mudancas"]["F1"]["fechamento"][0]
        self.assertEqual((item["depois"], item["tabela"]), (None, None))
        self.assertContains(r, "fechamento (removido)")
        self.assertContains(r, "Utilização\n\nFim.</pre>")

    def test_texto_escapado_no_diff_e_na_textarea(self):
        self.imprimir()
        self.nova_versao_f1(F1.replace("Fim.", "Fim <script>alert(1)</script> </textarea>."))
        r = self.abrir()
        corpo = r.content.decode()
        self.assertNotIn("<script>alert(1)", corpo)
        self.assertNotIn("</textarea>.", corpo)
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt; &lt;/textarea&gt;.", corpo)  # textarea
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;&nbsp;&lt;/textarea&gt;.", corpo)  # HtmlDiff

    def test_post_anonimo_vai_ao_login(self):
        self.client.logout()
        url = reverse("impresso_marcar", args=[self.rev.pk])
        r = self.client.post(url, {"hash_visto": "x"})
        self.assertRedirects(r, f"{reverse('admin:login')}?next={url}", fetch_redirect_response=False)
        self.assertFalse(FotoRevista.objects.exists())

    def planilha(self, nome):
        r = self.client.get(reverse("impresso_planilhas", args=[self.rev.pk]))
        self.assertEqual(r.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(zipfile.ZipFile(io.BytesIO(r.content)).read(nome)))

    def test_caractere_de_controle_nao_derruba_as_planilhas(self):
        Descritor.objects.filter(matriz_etapa_cc__etapa=self.e5, matriz_etapa_cc__componente=self.lp, codigo="D01").update(descricao="Des\x01critor\x0b 1")
        Projeto.objects.filter(pk=self.projeto_go.pk).update(adjetivo_patrio_plural_masculino="goia\x01nos\x0b")
        self.assertEqual(self.planilha("Projeto GO - LP 5EF.xlsx")["Níveis"]["C2"].value, "Descritor 1")
        self.assertEqual(self.planilha("Projeto GO - LP F2.xlsx").active["C3"].value, "Para os goianos.")

    def test_passo_do_projeto_sobrescreve_o_da_matriz(self):
        ProjetoPasso.objects.create(projeto=self.projeto_go, etapa=self.e5, componente=self.lp, ano=self.ano, passo=50)
        resolvido = resolver(carregar(self.rev.pk))["f1"]["dados"]["5ef"]
        vazias = {x["level"] for d in planilhas.descritores(resolvido) for ls in d[1]["scale"].values() for x in ls if x["content"] == "---"}
        self.assertTrue(vazias and all(n % 50 == 0 for n in vazias))  # o resolver usa 50
        self.assertEqual(self.planilha("Projeto GO - LP 5EF.xlsx")["Escala"]["B4"].value, 50)
        self.assertEqual(self.planilha("Projeto GO - LP 9EF.xlsx")["Escala"]["B4"].value, 25)

    def test_arquivada_sem_intervalo_minimo(self):
        Revista.objects.filter(pk=self.rev.pk).update(arquivada_em=timezone.now())
        self.assertIsNone(self.planilha("Projeto GO - LP 5EF.xlsx")["Escala"]["B4"].value)

    def test_baixar_planilhas(self):
        self.client.force_login(self.imp)
        r = self.client.get(reverse("impresso_planilhas", args=[self.rev.pk]))
        self.assertEqual(r["Content-Disposition"], 'attachment; filename="planilhas-go-lp-2026.zip"')
        z = zipfile.ZipFile(io.BytesIO(r.content))
        self.assertEqual(z.namelist(), ["Projeto GO - LP 5EF.xlsx", "Projeto GO - LP 9EF.xlsx", "Projeto GO - LP F2.xlsx"])
        escala = openpyxl.load_workbook(io.BytesIO(z.read("Projeto GO - LP 5EF.xlsx")))["Escala"]
        self.assertEqual([c.value for c in escala["B"]][:4], ["5EF", 0, 500, 25])  # passo da MEC
        self.assertEqual([c.value for c in escala[7]], ["Padrão 01", 0, 100, "#F72628"])
        self.assertEqual(escala.max_row, 11)  # 5 padrões (4 cortes)
        f2 = openpyxl.load_workbook(io.BytesIO(z.read("Projeto GO - LP F2.xlsx"))).active
        self.assertEqual([c.value for c in f2[3]], ["passo1", "Passo um", "Para os goianos."])
        # Ação da lista de revistas: redireciona para o mesmo download.
        r = self.client.post(reverse("admin:nucleo_revista_changelist"), {"action": "baixar_planilhas", "_selected_action": [self.rev.pk]})
        self.assertRedirects(r, reverse("impresso_planilhas", args=[self.rev.pk]), fetch_redirect_response=False)


def celulas(ws):
    """Valor, negrito, preenchimento, quebra e alinhamento de cada célula; larguras e congelamento."""
    grade = {  # o ExcelJS grava "" como célula de texto vazia; o openpyxl não grava a célula: mesmo efeito
        c.coordinate: (None if c.value == "" else c.value, bool(c.font.b), c.fill.fgColor.rgb if c.fill.fill_type else None, bool(c.alignment.wrap_text), c.alignment.vertical)
        for linha in ws.iter_rows()
        for c in linha
    }
    larguras = {i: d.width for d in ws.column_dimensions.values() if d.width for i in range(d.min, d.max + 1)}
    return grade, larguras, ws.freeze_panes


class GabaritoPlanilhaTest(SimpleTestCase):
    """Gabarito: o `service.ts` antigo rodado uma vez sobre revista-3cortes.json (5EF LP e 9EF MT)."""

    def test_celula_a_celula(self):
        revista = json.loads((FIXTURES / "revista-3cortes.json").read_text())
        for etapa, cc in (("5ef", "LP"), ("9ef", "MT")):
            gabarito = openpyxl.load_workbook(FIXTURES / f"planilha-3cortes-{cc}-{etapa}.xlsx")
            novo = openpyxl.load_workbook(io.BytesIO(planilhas.planilha_f1(revista[etapa], etapa, cc, CORES_ANTIGAS)))
            self.assertEqual(novo.sheetnames, gabarito.sheetnames)
            for aba in gabarito.sheetnames:
                with self.subTest(etapa=etapa, aba=aba):
                    self.assertEqual(celulas(novo[aba]), celulas(gabarito[aba]))

    def test_n_padroes(self):
        dados = {
            "scaleRange": {"min": 0, "max": 500}, "cortes": {f"padrao-{i}": 100 * i for i in range(1, 5)},
            "padroes": [{"chave": f"padrao-{k}", "rotulo": f"P{k}"} for k in range(1, 6)],
            "D01": {"description": "d", "prerequisites": [], "bncc": {"practices": "", "knowledge": "", "skills": ["(X) c", "Y"]},
                    "scale": {"padrao-5": [{"level": 450, "content": "alto"}], "padrao-1": [{"level": 50, "content": "---"}]}},
        }  # fmt: skip
        wb = openpyxl.load_workbook(io.BytesIO(planilhas.planilha_f1(dados, "5ef", "LP", {})))
        self.assertEqual([c.value for c in wb["Níveis"][1]][4:7], ["0 a 50", "50 a 450", "450 a 500"])
        self.assertEqual([c.value for c in wb["Níveis"][2]][4:], ["–", "–", "alto", None, None, "Y", "(X) c"])
        self.assertEqual([c.value for c in wb["Escala"]["A"]][5:], ["Padrão", "P1", "P2", "P3", "P4", "P5"])
        self.assertEqual(wb["Escala"]["C6"].value, "Fim (< / ≤ no 05)")
        self.assertEqual([c.value for c in wb["Escala"][11]][:3], ["P5", 400, 500])

    def test_formula_sai_como_texto(self):
        dados = json.loads((FIXTURES / "revista-3cortes.json").read_text())["9ef"]
        dados["D01"]["description"] = '=HYPERLINK("http://exemplo.invalid","x")'
        c = openpyxl.load_workbook(io.BytesIO(planilhas.planilha_f1(dados, "9ef", "MT", {})))["Níveis"]["C2"]
        self.assertEqual((c.value, c.data_type), ('=HYPERLINK("http://exemplo.invalid","x")', "s"))
