"""Cadastros iniciais e importação de matriz (S04).

Os 3 PDFs do SAEGO 2025 ficam em `revista/modelo/PDFs Matriz Comunitária/`, fora da imagem
do app. Os testes com PDF procuram os arquivos nas pastas de SAEGO_PDFS (separadas por ":";
padrão "/pdfs:/pdfs-alfa") e são pulados quando faltam. Exemplo:
  docker compose run --rm -v "$PWD/revista/modelo/PDFs Matriz Comunitária:/pdfs:ro" \
    app python manage.py test nucleo.tests.test_importacao
"""
import os
import tempfile
from datetime import timedelta
from io import StringIO
from pathlib import Path
from unittest import mock, skipUnless

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import IntegrityError
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook

from nucleo.importacao.pdf import tarefas
from nucleo.models import (
    Ano, Bloqueio, Componente, CorPadrao, Descritor, Etapa, Matriz, MatrizEtapaCC, Tarefa,
)
from nucleo.servicos import importar_matriz as importar_matriz_mod
from nucleo.servicos.importar_matriz import ImportacaoRecusada, importar_matriz

PASTAS_PDF = [Path(p) for p in os.environ.get("SAEGO_PDFS", "/pdfs:/pdfs-alfa").split(os.pathsep) if p]


def pdf_saego(nome):
    return next((p / nome for p in PASTAS_PDF if (p / nome).exists()), None)


PDF_ALFA = pdf_saego("SAEGO_2025_RE_Alfa_2EF_5EF Web.pdf")
PDF_LP = pdf_saego("SAEGO_2025_RE_LP Web.pdf")
PDF_MT = pdf_saego("SAEGO_2025_RE_MT Web.pdf")


def cadastrar(ano=2026):
    call_command("carregar_cadastros", "--ano", str(ano), stdout=StringIO())
    return Ano.objects.get(ano=ano)


class CarregarCadastrosTest(TestCase):
    def test_cria_tudo_e_e_idempotente(self):
        saida = StringIO()
        call_command("carregar_cadastros", "--ano", "2026", stdout=saida)
        self.assertIn("12 registro(s) criado(s)", saida.getvalue())  # ano + 4 + 2 + 5
        Etapa.objects.filter(codigo="5EF").update(escala_max=999)  # edição no admin
        CorPadrao.objects.filter(padrao=1).update(nome="Abaixo do básico")
        saida = StringIO()
        call_command("carregar_cadastros", "--ano", "2026", stdout=saida)
        self.assertIn("0 registro(s) criado(s)", saida.getvalue())
        self.assertEqual(
            list(Etapa.objects.order_by("ordem").values_list("codigo", "escala_max", "ordem")),
            [("2EF", 1000, 1), ("5EF", 999, 2), ("9EF", 500, 3), ("3EM", 500, 4)],
        )
        self.assertEqual(sorted(Componente.objects.values_list("codigo", flat=True)), ["LP", "MT"])
        self.assertEqual(
            list(CorPadrao.objects.order_by("padrao").values_list("cor", flat=True)),
            ["#F72628", "#FF8D43", "#B3CCFF", "#0ED3FF", "#FFFFFF"],
        )
        self.assertEqual(  # D-175: nome padrão por ano, sem desfazer a edição
            list(CorPadrao.objects.order_by("padrao").values_list("nome", flat=True)),
            ["Abaixo do básico", "Padrão 02", "Padrão 03", "Padrão 04", "Padrão 05"],
        )


def planilha(caminho, abas, bncc=("P", "O", "HR", "HC")):
    """Workbook no layout da exportação: abas = [(cc, etapa, intervalos, linhas)]."""
    wb = Workbook()
    wb.remove(wb.active)
    for cc, etapa, intervalos, linhas in abas:
        ws = wb.create_sheet("Níveis")
        inicio = ["Etapa", "Descritor" if cc == "LP" else "Cód. Descritor", "Descrição", "Pre requisito"]
        ws.append(inicio + [f"{a} a {b}" for a, b in intervalos] + list(bncc))
        for linha in linhas:
            ws.append([etapa, *linha])
    wb.create_sheet("Escala").append(["Etapa", "5EF"])
    wb.save(caminho)
    return caminho


class ImportarXlsxTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.ano = cadastrar()
        cls.matriz = Matriz.objects.create(ano=cls.ano, nome="Comunitária")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.xlsx = planilha(
            Path(self.tmp.name) / "m.xlsx",
            [
                ("LP", "5ef", [(0, 125), (125, 150), (150, 500)], [
                    ["D01", "Localizar informação.", "Ler#Escrever", "–", "Tarefa A#Tarefa B", "Tarefa C",
                     "Leitura", "Objeto", "EF15LP01", "(EF15LP02) Algo."],
                    ["D02", "Inferir.", "", "–", "---", "–", "", "", "", ""],
                ]),
                ("MT", "5EF", [(0, 200), (200, 500)], [
                    ["D01", "Contar.", "", "...", "Contar até 10", "Eixo", "Unidade", "", ""],
                ]),
            ],
        )

    def test_grava_descritores_e_tarefas(self):
        self.assertEqual(importar_matriz(self.matriz, self.xlsx), {("LP", "5EF"): 2, ("MT", "5EF"): 1})
        d = Descritor.objects.get(matriz_etapa_cc__componente__codigo="LP", codigo="D01")
        self.assertEqual(d.prerequisitos, ["Ler", "Escrever"])
        self.assertEqual((d.praticas, d.habilidade_correspondente), ("Leitura", "(EF15LP02) Algo."))
        self.assertEqual(
            list(d.tarefa_set.order_by("ordem").values_list("inicio", "fim", "texto")),
            [(150, 500, "Tarefa C"), (125, 150, "Tarefa A"), (125, 150, "Tarefa B")],
        )
        self.assertFalse(Tarefa.objects.filter(descritor__codigo="D02").exists())  # "–" e "---"
        self.assertEqual(d.history.count(), 1)  # bulk_create_with_history

    def test_reimportar_sem_substituir_e_recusado(self):
        importar_matriz(self.matriz, self.xlsx)
        with self.assertRaisesMessage(ImportacaoRecusada, "LP 5EF, MT 5EF"):
            importar_matriz(self.matriz, self.xlsx)

    def test_passo_omitido_em_mec_nova_e_nulo(self):
        importar_matriz(self.matriz, self.xlsx)
        self.assertEqual(set(MatrizEtapaCC.objects.values_list("passo", flat=True)), {None})

    def test_passo_nao_positivo_e_recusado(self):
        with self.assertRaisesMessage(ImportacaoRecusada, "passo deve ser positivo"):
            importar_matriz(self.matriz, self.xlsx, passo=0)

    def test_substituir_preserva_mec_e_passo_ou_usa_o_informado(self):
        importar_matriz(self.matriz, self.xlsx, passo=25)
        ids = set(MatrizEtapaCC.objects.values_list("id", flat=True))
        Tarefa.objects.update(texto="editada")
        importar_matriz(self.matriz, self.xlsx, substituir=True)
        self.assertEqual(set(MatrizEtapaCC.objects.values_list("id", flat=True)), ids)
        self.assertEqual(set(MatrizEtapaCC.objects.values_list("passo", flat=True)), {25})
        self.assertFalse(Tarefa.objects.filter(texto="editada").exists())
        self.assertEqual(Descritor.objects.count(), 3)
        self.assertEqual(
            Descritor.history.filter(history_type="+").first().history_change_reason, "importação de m.xlsx"
        )
        importar_matriz(self.matriz, self.xlsx, substituir=True, passo=None)
        self.assertEqual(set(MatrizEtapaCC.objects.values_list("passo", flat=True)), {None})

    def test_substituir_mt_nao_mexe_em_lp(self):
        importar_matriz(self.matriz, self.xlsx)
        Tarefa.objects.filter(descritor__matriz_etapa_cc__componente__codigo="LP").update(texto="editada")
        so_mt = planilha(Path(self.tmp.name) / "mt.xlsx", [
            ("MT", "5EF", [(0, 200), (200, 500)], [["D09", "Medir.", "", "...", "Medir", "", "", "", ""]]),
        ])
        self.assertEqual(importar_matriz(self.matriz, so_mt, substituir=True), {("MT", "5EF"): 1})
        self.assertEqual(Tarefa.objects.filter(texto="editada").count(), 3)
        self.assertEqual(
            sorted(Descritor.objects.values_list("matriz_etapa_cc__componente__codigo", "codigo")),
            [("LP", "D01"), ("LP", "D02"), ("MT", "D09")],
        )

    def test_falha_no_meio_da_substituicao_desfaz_tudo(self):
        importar_matriz(self.matriz, self.xlsx, passo=25)
        antes = sorted(Tarefa.objects.values_list("descritor__codigo", "texto"))
        original = importar_matriz_mod.bulk_create_with_history

        def falha_nas_tarefas(objs, model, **kw):
            if model is Tarefa:
                raise IntegrityError("simulada")
            return original(objs, model, **kw)

        with mock.patch.object(importar_matriz_mod, "bulk_create_with_history", falha_nas_tarefas):
            with self.assertRaises(IntegrityError):
                importar_matriz(self.matriz, self.xlsx, substituir=True, passo=40)
        self.assertEqual(sorted(Tarefa.objects.values_list("descritor__codigo", "texto")), antes)
        self.assertEqual(set(MatrizEtapaCC.objects.values_list("passo", flat=True)), {25})

    def test_bloqueio_ativo_recusa_substituir(self):
        importar_matriz(self.matriz, self.xlsx)
        usuario = get_user_model().objects.create(username="pedagogica")
        lp = MatrizEtapaCC.objects.get(componente__codigo="LP")
        trava = Bloqueio.objects.create(
            matriz_etapa_cc=lp, usuario=usuario, sessao_edicao="aba",
            expira_em=timezone.now() + timedelta(seconds=30),
        )
        with self.assertRaisesMessage(ImportacaoRecusada, "em edição agora: LP 5EF (pedagogica)"):
            importar_matriz(self.matriz, self.xlsx, substituir=True)
        trava.expira_em = timezone.now() - timedelta(seconds=1)  # vencida: não impede
        trava.save()
        importar_matriz(self.matriz, self.xlsx, substituir=True)
        self.assertTrue(Bloqueio.objects.filter(pk=trava.pk, matriz_etapa_cc=lp).exists())

    def test_planilha_invalida_e_recusada_antes_do_banco(self):
        casos = {
            "início ≥ fim": dict(abas=[("LP", "5EF", [(0, 250), (250, 250)], [["D01", "x", "", "a", "b"]])]),
            "não continua": dict(abas=[("LP", "5EF", [(0, 200), (250, 500)], [["D01", "x", "", "a", "b"]])]),
            "faltam colunas da BNCC": dict(abas=[("LP", "5EF", [(0, 500)], [["D01", "x", "", "a"]])], bncc=("P",)),
            "código D01 repetido": dict(abas=[("LP", "5EF", [(0, 500)], [["D01", "x", "", "a"], ["D01", "y", "", "b"]])]),
        }
        for msg, kw in casos.items():
            with self.subTest(msg), self.assertRaisesMessage(ValueError, msg):
                importar_matriz(self.matriz, planilha(Path(self.tmp.name) / "i.xlsx", **kw))
        corrompido = Path(self.tmp.name) / "c.xlsx"
        corrompido.write_bytes(b"isto nao e um zip")
        with self.assertRaisesMessage(ValueError, "arquivo não é um .xlsx válido"):
            importar_matriz(self.matriz, corrompido)
        self.assertFalse(MatrizEtapaCC.objects.exists())

    def test_bncc_com_espacos_normalizados(self):
        xlsx = planilha(Path(self.tmp.name) / "b.xlsx", [
            ("LP", "5EF", [(0, 500)], [["D01", "x", "", "a", " Leitura\n e  escuta ", "", "", ""]]),
        ])
        importar_matriz(self.matriz, xlsx)
        self.assertEqual(Descritor.objects.get().praticas, "Leitura e escuta")

    def test_escala_diferente_da_etapa_e_recusada(self):
        xlsx = planilha(Path(self.tmp.name) / "e.xlsx", [("LP", "2EF", [(0, 500)], [["D01", "x", "", "y"]])])
        with self.assertRaisesMessage(ImportacaoRecusada, "vai até 500, a etapa até 1000"):
            importar_matriz(self.matriz, xlsx)
        self.assertFalse(MatrizEtapaCC.objects.exists())

    def test_etapa_nao_cadastrada_e_recusada(self):
        xlsx = planilha(Path(self.tmp.name) / "e.xlsx", [("LP", "1EF", [(0, 500)], [["D01", "x", "", "y"]])])
        with self.assertRaisesMessage(ImportacaoRecusada, "etapa 1EF"):
            importar_matriz(self.matriz, xlsx)

    def test_comando_cria_matriz_e_recusa_reimportar(self):
        saida = StringIO()
        call_command("importar_matriz", str(self.xlsx), "--ano", "2026", "--nome", "Outra", "--passo", "25",
                     stdout=saida)
        self.assertIn("Outra (2026) LP 5EF: 2 descritores", saida.getvalue())
        self.assertEqual(MatrizEtapaCC.objects.get(matriz__nome="Outra", componente__codigo="MT").passo, 25)
        with self.assertRaisesMessage(CommandError, "confirme com substituir"):
            call_command("importar_matriz", str(self.xlsx), "--ano", "2026", "--nome", "Outra")

    def test_comando_nao_deixa_matriz_vazia_quando_recusa(self):
        with self.assertRaises(CommandError):
            call_command("importar_matriz", "x.docx", "--ano", "2026", "--nome", "Nova")
        escala_errada = planilha(Path(self.tmp.name) / "e.xlsx", [("LP", "2EF", [(0, 500)], [["D01", "x", "", "y"]])])
        with self.assertRaisesMessage(CommandError, "a etapa até 1000"):
            call_command("importar_matriz", str(escala_errada), "--ano", "2026", "--nome", "Nova")
        self.assertFalse(Matriz.objects.filter(nome="Nova").exists())


class TarefasDoPdfTest(TestCase):
    def test_linhas_vazias_nao_viram_tarefa_e_ordem_e_de_cima_para_baixo(self):
        niveis = [
            {"level": 500, "fim": 600, "content": "B"},
            {"level": 600, "fim": 1000, "content": "---"},
            {"level": 600, "fim": 1000, "content": "A"},
        ]
        self.assertEqual(
            tarefas(niveis),
            [{"inicio": 600, "fim": 1000, "texto": "A"}, {"inicio": 500, "fim": 600, "texto": "B"}],
        )


class PdfSaegoTest(TestCase):
    """Importa um PDF inteiro (descritores contados com o extrator antigo, scripts/)."""

    PDF = None
    ESPERADO = {}

    @classmethod
    def setUpTestData(cls):
        cls.matriz = Matriz.objects.create(ano=cadastrar(), nome="Comunitária")
        cls.resumo = importar_matriz(cls.matriz, cls.PDF)

    def desc(self, cc, etapa, codigo):
        return Descritor.objects.get(
            matriz_etapa_cc__componente__codigo=cc, matriz_etapa_cc__etapa__codigo=etapa, codigo=codigo
        )

    def tarefas(self, d):
        return list(d.tarefa_set.order_by("ordem").values_list("inicio", "fim", "texto"))

    def test_contagens_iguais_ao_extrator_antigo(self):
        self.assertEqual(self.resumo, self.ESPERADO)
        self.assertFalse(Tarefa.objects.filter(texto="---").exists())


@skipUnless(PDF_ALFA, f"PDF Alfa não encontrado em {PASTAS_PDF}")
class PdfAlfaTest(PdfSaegoTest):
    PDF = PDF_ALFA
    ESPERADO = {("LP", "2EF"): 8, ("LP", "5EF"): 15, ("MT", "2EF"): 33, ("MT", "5EF"): 28}

    def test_lp_2ef_d01(self):
        d = self.desc("LP", "2EF", "D01")
        self.assertEqual(d.descricao, "Relacionar elementos sonoros das palavras com sua representação escrita.")
        self.assertEqual(len(d.prerequisitos), 13)
        self.assertEqual(d.prerequisitos[6], "Reconhecer letras do alfabeto.")  # fim da coluna esquerda
        ts = self.tarefas(d)
        self.assertLessEqual({t[0] for t in ts}, {782, 732, 675, 616, 556, 503, 446, 390})
        self.assertEqual([t[0] for t in ts], sorted((t[0] for t in ts), reverse=True))
        self.assertTrue(ts[0][2].startswith('Reconhecer o som da letra "C" inicial'))
        self.assertTrue(any("/ʒ/ e /g/" in t[2] for t in ts))  # glifo deslocado no PDF

    def test_lp_2ef_d02_bncc(self):
        self.assertTrue(self.desc("LP", "2EF", "D02").habilidade_correspondente.startswith("(EF12LP01)"))

    def test_mt_2ef_d06_rotulos_trocados(self):
        d = self.desc("MT", "2EF", "D06")
        self.assertEqual(d.habilidades_relacionadas, "EF02MA01")
        self.assertTrue(d.descricao.endswith("(ou valor relativo) em um número natural de até 3 ordens."))

    def test_mt_2ef_d27_tarefas_sem_linha_vazia(self):
        faixas = [(i, f, t.split(",")[0]) for i, f, t in self.tarefas(self.desc("MT", "2EF", "D27")) if i >= 522]
        self.assertEqual(faixas, [
            (681, 1000, "Determinar a data de início de um acontecimento"),
            (681, 1000, "Determinar a data de término de um acontecimento"),
            (635, 681, "Determinar a duração de um evento (em dias) dadas as informações sobre a data de início e de término"),
            # (596, 635, "---") do PDF não vira Tarefa (D-091)
            (530, 596, "Identificar o dia da semana a partir de um contexto do cotidiano"),
            (530, 596, "Determinar o dia do mês em que uma pessoa retornará de uma viagem"),
            (530, 596, "Determinar o mês em que termina um acontecimento"),
            (522, 530, "Identificar"),
        ])


@skipUnless(PDF_LP, f"PDF LP não encontrado em {PASTAS_PDF}")
class PdfLpTest(PdfSaegoTest):
    PDF = PDF_LP
    ESPERADO = {("LP", "9EF"): 21, ("LP", "3EM"): 21}

    def test_bncc(self):
        # Valor longo: a 1ª linha ficava no rótulo de cima ("relacionadas").
        d = self.desc("LP", "3EM", "D05")
        self.assertEqual(d.habilidades_relacionadas, "EF69LP33, EF07LP02, EM13LP18, EM13LP19")
        self.assertTrue(d.habilidade_correspondente.startswith("(EM13LP15) Planejar"))
        # Coluna de valores em x=154 (< 165): o valor era lido como rótulo.
        d = self.desc("LP", "9EF", "D03")
        self.assertEqual(d.habilidades_relacionadas, "EF69LP20")
        self.assertTrue(d.habilidade_correspondente.startswith("(EF69LP47)"))


@skipUnless(PDF_MT, f"PDF MT não encontrado em {PASTAS_PDF}")
class PdfMtTest(PdfSaegoTest):
    PDF = PDF_MT
    ESPERADO = {("MT", "9EF"): 37, ("MT", "3EM"): 35}


del PdfSaegoTest  # base abstrata: sem PDF próprio
