"""Admin de cadastros, matriz e cortes (S11a): testes pelo client."""

import io
import tempfile
import threading
from unittest import mock
from pathlib import Path

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import Client, TestCase, TransactionTestCase
from django.urls import reverse
from openpyxl import load_workbook

from nucleo.admin.cortes import numerar_grupos
from nucleo.models import (
    Ano, Componente, CorPadrao, Corte, Descritor, Etapa, Matriz, MatrizEtapaCC, Parceiro,
    ParceiroAno, Projeto, Revista, Tarefa,
)  # fmt: skip
from nucleo.servicos.cascata import gravar_cortes
from nucleo.tests.test_importacao import cadastrar, planilha
from nucleo.tests.test_resolver import usuario

# Planilha "Cortes dos Padrões por Etapa e Parceiro.xlsx" (2025), com os grupos corrigidos
# (D-049): 9EF-MT SE 6 e SALVADOR 7; 5EF-MT SALVADOR 6. A aba 2EF-LP também tinha erro
# (MA/PI/SALVADOR em grupo separado do AM, mesmos cortes): aqui já corrigida.
GRUPOS = {
    ("9EF", "MT"): [
        (1, (199, 274, 324), ["MS"]),
        (2, (200, 250, 325), ["MT"]),
        (3, (225, 245, 280), ["PE"]),
        (4, (225, 275, 325), ["CE", "GO", "PI"]),
        (5, (225, 300, 350), ["BA", "ES", "Florianópolis", "MA", "MG", "PB", "PR", "RN", "RO", "JOINVILLE", "PORTO ALEGRE"]),
        (6, (250, 300, 350), ["SE"]),
        (7, (270, 305, 323), ["SALVADOR"]),
    ],
    ("5EF", "MT"): [
        (1, (124, 174, 224), ["MS"]),
        (2, (150, 185, 220), ["PE"]),
        (3, (150, 200, 250), ["CE", "GO"]),
        (4, (175, 225, 275), ["AC", "AP", "BA", "DF", "ES", "Florianópolis", "MA", "MG", "PB", "PI",
                              "PORTO ALEGRE", "PR", "RN", "RO", "SE", "TO"]),
        (5, (175, 225, 300), ["MT"]),
        (6, (200, 235, 253), ["SALVADOR"]),
    ],
    ("2EF", "LP"): [
        (1, (350, 400, 500), ["GO"]),
        (2, (350, 450, 525), ["PE"]),
        (3, (350, 450, 600), ["MG"]),
        (4, (399, 499, 599), ["MS"]),
        (5, (400, 500, 600), ["MT"]),
        (6, (500, 600, 700), ["AC", "AL", "AP", "BA", "DF", "ES", "CE", "Florianópolis", "PA", "PB", "PR",
                              "RN", "RJ", "RO", "TO", "JOINVILLE", "PORTO ALEGRE", "RR", "SC"]),
        (7, (518, 600, 667), ["AM", "MA", "PI", "SALVADOR"]),
        (8, (725, 800), ["SE"]),
    ],
}  # fmt: skip


def dados_da_grade(resp):
    dados = {}
    for linha in resp.context["linhas"]:
        dados[f"o-{linha['parceiro'].pk}"] = linha["original"]
        for cel in linha["celulas"]:
            dados[f"c-{cel['chave']}"] = cel["texto"]
    return dados


def dados_do_form(resp):
    """POST equivalente ao change form renderizado (form principal, inlines e sessão do bloqueio)."""
    forms = [resp.context["adminform"].form]
    dados = {}
    for inline in resp.context["inline_admin_formsets"]:
        forms += [inline.formset.management_form, *inline.formset.forms]
    for form in forms:
        for nome in form.fields:
            valor = form[nome].value()
            if valor in (None, False):
                continue
            dados[form.add_prefix(nome)] = "\n".join(valor) if isinstance(valor, list) else valor
    if resp.context.get("bloqueio"):
        dados["_bloqueio_sessao"] = resp.context["bloqueio"]["sessao"]
    return dados


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.a26 = cadastrar(2026)
        cls.e = {e.codigo: e for e in Etapa.objects.filter(ano=cls.a26)}
        cls.c = {c.codigo: c for c in Componente.objects.filter(ano=cls.a26)}
        cls.matriz = Matriz.objects.create(ano=cls.a26, nome="Comunitária")
        cls.mec = MatrizEtapaCC.objects.create(
            matriz=cls.matriz, etapa=cls.e["5EF"], componente=cls.c["LP"], ano=cls.a26, passo=25
        )
        for i in (1, 2):
            d = Descritor.objects.create(
                matriz_etapa_cc=cls.mec, codigo=f"D0{i}", descricao="Lê, interpreta.", ordem=i,
                prerequisitos=["Ler, escrever", "Contar"],
            )  # fmt: skip
            Tarefa.objects.create(descritor=d, inicio=100, fim=125, texto="t", ordem=1)
        cls.root = usuario("root", is_superuser=True)

    def setUp(self):
        self.client.force_login(self.root)


class MatrizTest(Base):
    def test_duplicar_para_outro_ano_com_diff(self):
        a27 = cadastrar(2027)
        url = reverse("admin:nucleo_matriz_duplicar", args=[self.matriz.pk])
        resp = self.client.post(url, {"ano": a27.pk, "nome": "Exceção GO"})
        nova = Matriz.objects.get(ano=a27, nome="Exceção GO")
        self.assertRedirects(resp, reverse("admin:nucleo_matriz_change", args=[nova.pk]))
        self.assertEqual(nova.duplicada_de, self.matriz)
        mec = MatrizEtapaCC.objects.get(matriz=nova)
        self.assertEqual((mec.etapa.ano, mec.componente.codigo, mec.passo), (a27, "LP", 25))
        self.assertEqual(Descritor.objects.filter(matriz_etapa_cc=mec).count(), 2)
        self.assertEqual(Tarefa.objects.filter(descritor__matriz_etapa_cc=mec).count(), 2)
        self.assertEqual(Descritor.objects.filter(matriz_etapa_cc=self.mec).count(), 2)  # origem intacta

        diff = reverse("admin:nucleo_matriz_diferencas", args=[nova.pk])
        self.assertContains(self.client.get(diff), "Nenhuma diferença")
        Descritor.objects.filter(matriz_etapa_cc=mec, codigo="D02").update(descricao="Infere.")
        mec.passo = 50
        mec.save()
        resp = self.client.get(diff)
        self.assertContains(resp, "intervalo mínimo")
        self.assertContains(resp, "5EF LP D02")
        self.assertContains(resp, "Infere.")

    def test_duplicar_recusa_ano_sem_etapa_e_nome_repetido(self):
        a27 = Ano.objects.create(ano=2027)
        url = reverse("admin:nucleo_matriz_duplicar", args=[self.matriz.pk])
        self.assertContains(self.client.post(url, {"ano": a27.pk, "nome": "X"}), "não cadastrados em 2027")
        self.assertContains(
            self.client.post(url, {"ano": self.a26.pk, "nome": "Comunitária"}), "já existe a matriz"
        )
        self.assertEqual(Matriz.objects.count(), 1)

    def test_reimportar_pede_nome_da_matriz_para_substituir(self):
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = planilha(Path(tmp) / "m.xlsx", [("LP", "5EF", [(0, 200), (200, 500)], [
                ["D09", "Novo.", "", "–", "Tarefa N", "", "", "", ""],
            ])]).read_bytes()  # fmt: skip
        url = reverse("admin:nucleo_matriz_reimportar", args=[self.matriz.pk])

        def enviar(**extra):
            return self.client.post(url, {"arquivo": SimpleUploadedFile("m.xlsx", xlsx), **extra})

        self.assertContains(enviar(), "Digite Comunitária para confirmar")
        self.assertContains(enviar(confirmacao="comunitaria"), "Digite Comunitária")
        self.assertEqual(set(self.mec.descritor_set.values_list("codigo", flat=True)), {"D01", "D02"})
        self.assertRedirects(enviar(confirmacao="Comunitária"), reverse("admin:nucleo_matriz_change", args=[self.matriz.pk]))
        self.assertEqual(list(self.mec.descritor_set.values_list("codigo", flat=True)), ["D09"])
        self.mec.refresh_from_db()
        self.assertEqual(self.mec.passo, 25)  # passo vazio mantém

    def test_reimportar_arquivo_invalido(self):
        url = reverse("admin:nucleo_matriz_reimportar", args=[self.matriz.pk])
        resp = self.client.post(url, {"arquivo": SimpleUploadedFile("m.txt", b"%PDF-1.4")})
        self.assertContains(resp, "formato não suportado")


class ImpactoTest(Base):
    """D-076: antes de salvar MEC/Descritor, confirmação com o nº de revistas afetadas."""

    def com_revista(self):
        pa = ParceiroAno.objects.create(parceiro=Parceiro.objects.create(nome="GO"), ano=self.a26, matriz=self.matriz)
        gravar_cortes(pa, [(self.e["5EF"], self.c["LP"], [150, 200])], self.root)
        # Outro parceiro com corte só em MT: a revista de MT não muda com a MEC de LP.
        pa2 = ParceiroAno.objects.create(parceiro=Parceiro.objects.create(nome="PE"), ano=self.a26, matriz=self.matriz)
        gravar_cortes(pa2, [(self.e["5EF"], self.c["MT"], [150, 200])], self.root)
        self.assertEqual(Revista.objects.count(), 2)

    def test_mec_pede_confirmacao_quando_afeta_revistas(self):
        self.com_revista()
        url = reverse("admin:nucleo_matrizetapacc_change", args=[self.mec.pk])
        resp = self.client.get(url)
        self.assertContains(resp, "afeta 1 revista(s) ativa(s)")
        dados = {**dados_do_form(resp), "passo": "50"}
        resp = self.client.post(url, dados)
        self.assertContains(resp, "Esta edição afeta 1 revista(s)")
        self.mec.refresh_from_db()
        self.assertEqual(self.mec.passo, 25)
        resp = self.client.post(url, {**dados, "confirmar_impacto": "on"})
        self.assertEqual(resp.status_code, 302)
        self.mec.refresh_from_db()
        self.assertEqual(self.mec.passo, 50)

    def test_descritor_com_tarefa_inline_e_prerequisito_com_virgula(self):
        self.com_revista()
        d = Descritor.objects.get(matriz_etapa_cc=self.mec, codigo="D01")
        url = reverse("admin:nucleo_descritor_change", args=[d.pk])
        dados = dados_do_form(self.client.get(url))
        dados["tarefa_set-0-texto"] = "nova"
        self.assertContains(self.client.post(url, dados), "Esta edição afeta 1")
        resp = self.client.post(url, {**dados, "confirmar_impacto": "on"})
        self.assertEqual(resp.status_code, 302)
        d.refresh_from_db()
        self.assertEqual(d.prerequisitos, ["Ler, escrever", "Contar"])
        self.assertEqual(d.tarefa_set.get().texto, "nova")

    def test_sem_revista_salva_direto(self):
        url = reverse("admin:nucleo_matrizetapacc_change", args=[self.mec.pk])
        dados = {**dados_do_form(self.client.get(url)), "passo": ""}
        self.assertEqual(self.client.post(url, dados).status_code, 302)
        self.mec.refresh_from_db()
        self.assertIsNone(self.mec.passo)

    def test_revista_arquivada_nao_conta(self):
        self.com_revista()
        Revista.objects.update(arquivada_em="2026-01-01T00:00Z")
        resp = self.client.get(reverse("admin:nucleo_matrizetapacc_change", args=[self.mec.pk]))
        self.assertNotContains(resp, "afeta")


class CortesTest(Base):
    url = reverse("admin:nucleo_corte_changelist")

    def celula(self, parceiro, etapa, cc):
        return f"{parceiro.pk}-{self.e[etapa].pk}-{self.c[cc].pk}"

    def grade(self, resp=None):
        """POST com o estado da grade renderizada (o que está nas células e o original da linha)."""
        return dados_da_grade(resp or self.client.get(self.url, {"ano": 2026}))

    def test_grade_cria_parceiro_ano_e_dispara_a_cascata(self):
        go, pe = Parceiro.objects.create(nome="GO"), Parceiro.objects.create(nome="PE")
        dados = self.grade()
        dados[f"c-{self.celula(go, '5EF', 'LP')}"] = "125, 175, 225"
        dados[f"c-{self.celula(go, '9EF', 'MT')}"] = "225 275 325"
        resp = self.client.post(self.url + "?ano=2026", dados)
        self.assertRedirects(resp, self.url + "?ano=2026")
        pa = ParceiroAno.objects.get(parceiro=go, ano=self.a26)
        self.assertEqual(
            sorted(Corte.objects.filter(parceiro_ano=pa).values_list("etapa__codigo", "componente__codigo", "valores")),
            [("5EF", "LP", [125, 175, 225]), ("9EF", "MT", [225, 275, 325])],
        )
        projeto = Projeto.objects.get(parceiro_ano=pa)
        self.assertEqual(set(Revista.objects.filter(projeto=projeto).values_list("componente__codigo", flat=True)), {"LP", "MT"})
        self.assertFalse(ParceiroAno.objects.filter(parceiro=pe).exists())  # sem corte, sem ParceiroAno

        # Apagar a célula de MT: a cascata arquiva a revista de MT.
        dados = self.grade()
        dados[f"c-{self.celula(go, '9EF', 'MT')}"] = ""
        self.client.post(self.url + "?ano=2026", dados)
        self.assertEqual(Corte.objects.filter(parceiro_ano=pa).count(), 1)
        self.assertIsNotNone(Revista.objects.get(projeto=projeto, componente__codigo="MT").arquivada_em)

    def test_erros_da_d055_aparecem_e_nada_grava(self):
        go, pe = Parceiro.objects.create(nome="GO"), Parceiro.objects.create(nome="PE")
        dados = self.grade()
        dados[f"c-{self.celula(go, '5EF', 'LP')}"] = "200, 150"
        dados[f"c-{self.celula(go, '5EF', 'MT')}"] = "100, 500"
        dados[f"c-{self.celula(pe, '2EF', 'LP')}"] = "abc"
        dados[f"c-{self.celula(pe, '9EF', 'LP')}"] = "100, 200"  # válida, mas não grava
        resp = self.client.post(self.url + "?ano=2026", dados)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "5EF LP: cortes precisam ser crescentes.")
        self.assertContains(resp, "5EF MT: cortes precisam estar entre 1 e 499.")
        self.assertContains(resp, "2EF LP: cortes precisam ser inteiros")
        self.assertContains(resp, 'value="abc"')
        self.assertFalse(Corte.objects.exists())
        self.assertFalse(ParceiroAno.objects.exists())

    def test_conflito_com_edicao_de_outra_aba(self):
        go = Parceiro.objects.create(nome="GO")
        dados = self.grade()
        pa = ParceiroAno.objects.create(parceiro=go, ano=self.a26)
        gravar_cortes(pa, [(self.e["5EF"], self.c["LP"], [100])], self.root)  # outra aba gravou
        dados[f"c-{self.celula(go, '5EF', 'MT')}"] = "150"
        resp = self.client.post(self.url + "?ano=2026", dados)
        self.assertContains(resp, "mudaram desde que a página foi aberta (agora: ")
        self.assertContains(resp, 'value="150"')  # o que o usuário digitou continua na tela
        self.assertEqual(Corte.objects.get().valores, [100])
        # Salvar de novo a mesma tela sobrescreve de propósito (a linha "atual" virou o original).
        self.assertEqual(self.client.post(self.url + "?ano=2026", self.grade(resp)).status_code, 302)
        self.assertEqual(Corte.objects.get().valores, [150])

    def test_celula_ausente_e_erro_e_nao_apagar(self):
        go = Parceiro.objects.create(nome="GO")
        pa = ParceiroAno.objects.create(parceiro=go, ano=self.a26)
        gravar_cortes(pa, [(self.e["5EF"], self.c["LP"], [100])], self.root)
        dados = self.grade()
        del dados[f"c-{self.celula(go, '5EF', 'LP')}"]
        self.assertContains(self.client.post(self.url + "?ano=2026", dados), "célula ausente no envio")
        self.assertEqual(Corte.objects.get().valores, [100])

    def test_70_parceiros_cabem_no_limite_de_campos_do_post(self):
        Parceiro.objects.bulk_create(Parceiro(nome=f"P{i:02}") for i in range(70))
        dados = self.grade()
        self.assertLess(len(dados), 1000)  # DATA_UPLOAD_MAX_NUMBER_FIELDS
        p = Parceiro.objects.get(nome="P69")
        dados[f"c-{self.celula(p, '3EM', 'MT')}"] = "250, 300"
        self.assertEqual(self.client.post(self.url + "?ano=2026", dados).status_code, 302)
        self.assertEqual(Corte.objects.get().parceiro_ano.parceiro, p)


class GruposTest(Base):
    def test_numeracao_d056(self):
        for aba, grupos in GRUPOS.items():
            calculado = numerar_grupos([v for _, v, ps in grupos for _ in ps])
            self.assertEqual({v: g for g, v, _ in grupos}, calculado, aba)

    def test_planilha_exportada(self):
        for _, _, parceiros in GRUPOS[("2EF", "LP")]:
            for nome in parceiros:
                Parceiro.objects.get_or_create(nome=nome)
        for (etapa, cc), grupos in GRUPOS.items():
            for _, valores, parceiros in grupos:
                for nome in parceiros:
                    pa, _ = ParceiroAno.objects.get_or_create(parceiro=Parceiro.objects.get_or_create(nome=nome)[0], ano=self.a26)
                    Corte.objects.create(parceiro_ano=pa, etapa=self.e[etapa], componente=self.c[cc], ano=self.a26, valores=list(valores))
        resp = self.client.get(reverse("admin:nucleo_corte_exportar"), {"ano": 2026})
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, [f"{e} - {c}" for e in ("2EF", "5EF", "9EF", "3EM") for c in ("LP", "MT")])
        for (etapa, cc), grupos in GRUPOS.items():
            linhas = list(wb[f"{etapa} - {cc}"].iter_rows(values_only=True))
            self.assertEqual(linhas[0], ("Grupo", "Parceiro", "Etapa", cc))
            obtido, grupo = {}, None
            for g, parceiro, e, valores in linhas[1:]:
                grupo = g or grupo  # nº só na 1ª linha do grupo
                self.assertEqual(e, etapa)
                obtido[parceiro] = (grupo, valores)
            esperado = {p: (g, ", ".join(map(str, v))) for g, v, ps in grupos for p in ps}
            self.assertEqual(obtido, esperado)
        aba = {p: g for g, p, *_ in wb["9EF - MT"].iter_rows(values_only=True)}
        self.assertTrue(all(ws.title for ws in wb.worksheets))
        self.assertEqual((aba["SE"], aba["SALVADOR"]), (6, 7))
        self.assertEqual({p: g for g, p, *_ in wb["5EF - MT"].iter_rows(values_only=True)}["SALVADOR"], 6)


class PlanilhaSeguraTest(Base):
    def test_texto_que_parece_formula_fica_texto(self):
        pa = ParceiroAno.objects.create(parceiro=Parceiro.objects.create(nome="=1+1"), ano=self.a26)
        Corte.objects.create(parceiro_ano=pa, etapa=self.e["5EF"], componente=self.c["LP"], ano=self.a26, valores=[100])
        Componente.objects.filter(pk=self.c["MT"].pk).update(codigo="M/T:[x]")
        resp = self.client.get(reverse("admin:nucleo_corte_exportar"), {"ano": 2026})
        wb = load_workbook(io.BytesIO(resp.content))
        celula = wb["5EF - LP"]["B2"]
        self.assertEqual((celula.value, celula.data_type), ("=1+1", "s"))
        self.assertIn("5EF - M-T--x-", wb.sheetnames)



class SemAnoTest(TestCase):
    def test_exportar_sem_ano_404(self):
        self.client.force_login(usuario("root", is_superuser=True))
        self.assertEqual(self.client.get(reverse("admin:nucleo_corte_exportar")).status_code, 404)
        self.assertContains(self.client.get(reverse("admin:nucleo_corte_changelist")), "Nenhum ano aberto")


class CadastroTest(Base):
    def test_paginas_abrem(self):
        d = Descritor.objects.first()
        for nome, args in (
            ("nucleo_ano_changelist", []), ("nucleo_parceiro_add", []), ("nucleo_etapa_changelist", []),
            ("nucleo_componente_changelist", []), ("nucleo_corpadrao_changelist", []),
            ("nucleo_matriz_changelist", []), ("nucleo_matriz_add", []), ("nucleo_matriz_change", [self.matriz.pk]),
            ("nucleo_matrizetapacc_changelist", []), ("nucleo_descritor_changelist", []),
            ("nucleo_descritor_change", [d.pk]), ("nucleo_matriz_reimportar", [self.matriz.pk]),
            ("auth_user_changelist", []), ("auth_user_add", []), ("index", []),
        ):  # fmt: skip
            self.assertEqual(self.client.get(reverse(f"admin:{nome}", args=args)).status_code, 200, nome)

    def test_cor_invalida_recusada_no_form(self):
        cor = CorPadrao.objects.get(ano=self.a26, padrao=1)
        url = reverse("admin:nucleo_corpadrao_change", args=[cor.pk])
        self.assertContains(self.client.post(url, {"padrao": 1, "cor": "red"}), "#RRGGBB")
        self.assertRedirects(self.client.post(url, {"padrao": 1, "cor": "#00ff00"}), reverse("admin:nucleo_corpadrao_changelist"))

    def test_abrir_ano(self):
        url = reverse("admin:nucleo_ano_abrir")
        resp = self.client.get(url)
        self.assertContains(resp, "O próximo que pode ser aberto é 2027")
        self.assertNotContains(resp, 'value="2027"')  # o usuário digita (não há desfazer)
        self.assertContains(self.client.post(url, {"ano": 2029}), "próximo ano que pode ser aberto é 2027")
        self.assertRedirects(self.client.post(url, {"ano": 2027}), reverse("admin:nucleo_ano_changelist"))
        self.assertTrue(Matriz.objects.filter(ano__ano=2027, duplicada_de=self.matriz).exists())


class PerfisTest(Base):
    """Permissões por perfil nas telas do S11a (D-079, D-136, D-142)."""

    def entrar(self, papel):
        self.client.force_login(usuario(papel, papel))

    def test_impresso_so_ve(self):
        self.entrar("Impresso")
        for nome, args in (
            ("nucleo_corte_changelist", []), ("nucleo_matriz_change", [self.matriz.pk]),
            ("nucleo_matrizetapacc_change", [self.mec.pk]), ("nucleo_ano_changelist", []),
            ("nucleo_corte_exportar", []), ("nucleo_matriz_diferencas", [self.matriz.pk]),
        ):  # fmt: skip
            self.assertEqual(self.client.get(reverse(f"admin:{nome}", args=args)).status_code, 200, nome)
        for nome, args in (
            ("nucleo_ano_abrir", []), ("nucleo_matriz_duplicar", [self.matriz.pk]),
            ("nucleo_matriz_reimportar", [self.matriz.pk]), ("auth_user_changelist", []),
        ):
            self.assertEqual(self.client.get(reverse(f"admin:{nome}", args=args)).status_code, 403, nome)
        self.assertEqual(self.client.post(reverse("admin:nucleo_corte_changelist"), {}).status_code, 403)
        resp = self.client.get(reverse("admin:nucleo_matrizetapacc_change", args=[self.mec.pk]))
        self.assertNotContains(resp, 'name="passo"')
        self.assertNotContains(self.client.get(reverse("admin:nucleo_ano_changelist")), "Abrir ano")

    def test_pedagogica_abre_ano_e_edita_matriz_mas_nao_usuarios(self):
        self.entrar("Pedagógica")
        self.assertContains(self.client.get(reverse("admin:nucleo_ano_changelist")), "Abrir ano")
        self.client.post(reverse("admin:nucleo_ano_abrir"), {"ano": 2027})
        self.assertTrue(Ano.objects.filter(ano=2027).exists())
        self.assertEqual(self.client.get(reverse("admin:nucleo_ano_add")).status_code, 403)
        self.assertEqual(self.client.get(reverse("admin:nucleo_matriz_duplicar", args=[self.matriz.pk])).status_code, 200)
        self.assertContains(self.client.get(reverse("admin:nucleo_matrizetapacc_change", args=[self.mec.pk])), 'name="passo"')
        self.assertContains(self.client.get(reverse("admin:nucleo_corte_changelist")), "Salvar cortes")
        self.assertEqual(self.client.get(reverse("admin:auth_user_changelist")).status_code, 403)

    def form_de(self, alvo):
        return self.client.get(reverse("admin:auth_user_change", args=[alvo.pk])).context["adminform"].form

    def post_usuario(self, alvo, **campos):
        dados = {"username": alvo.username, "is_active": "on", **campos}
        return self.client.post(reverse("admin:auth_user_change", args=[alvo.pk]), dados)

    def test_admin_gerencia_usuarios_sem_superusuario(self):
        self.entrar("Admin")
        lista = self.client.get(reverse("admin:auth_user_changelist"))
        self.assertNotContains(lista, ">root<")  # superusuário fora da lista
        self.assertEqual(self.client.get(reverse("admin:auth_user_change", args=[self.root.pk])).status_code, 302)
        senha_root = reverse("admin:auth_user_password_change", args=[self.root.pk])
        self.assertNotEqual(self.client.get(senha_root).status_code, 200)
        self.assertNotEqual(self.client.post(senha_root, {"password1": "x-nova-123!", "password2": "x-nova-123!"}).status_code, 302)
        self.root.refresh_from_db()
        self.assertTrue(self.root.check_password("x"))

        ped = usuario("Pedagógica", "Pedagógica")
        form = self.form_de(ped)
        for campo in ("is_superuser", "is_staff", "user_permissions"):
            self.assertNotIn(campo, form.fields)
        Group.objects.create(name="Outro")
        self.assertEqual(set(form.fields["groups"].queryset.values_list("name", flat=True)), {"Pedagógica", "Mídias Digitais"})
        resp = self.post_usuario(ped, groups=[Group.objects.get(name="Impresso").pk])
        self.assertEqual(resp.status_code, 200)  # escolha fora da lista: erro no form
        self.assertEqual(list(ped.groups.values_list("name", flat=True)), ["Pedagógica"])
        self.post_usuario(ped, groups=[Group.objects.get(name="Mídias Digitais").pk])
        self.assertEqual(list(ped.groups.values_list("name", flat=True)), ["Mídias Digitais"])

        resp = self.client.post(
            reverse("admin:auth_user_add"),
            {"username": "nova", "password1": "senha-longa-123", "password2": "senha-longa-123", "usable_password": "true"},
        )
        nova = User.objects.get(username="nova")
        self.assertTrue(nova.is_staff)
        self.assertFalse(nova.is_superuser)

    def test_admin_nao_se_poe_no_impresso_nem_mexe_em_outro_admin(self):
        eu = usuario("Admin", "Admin")
        self.client.force_login(eu)
        self.assertNotIn("groups", self.form_de(eu).fields)  # readonly para si mesmo
        self.assertNotIn("is_active", self.form_de(eu).fields)
        self.post_usuario(eu, groups=[Group.objects.get(name="Impresso").pk], is_active="")
        eu.refresh_from_db()
        self.assertEqual(list(eu.groups.values_list("name", flat=True)), ["Admin"])
        self.assertTrue(eu.is_active)

        outro = User.objects.create_user("outro-admin", is_staff=True)
        outro.groups.add(Group.objects.get(name="Admin"))
        self.assertEqual(self.post_usuario(outro, is_active="").status_code, 403)
        outro.refresh_from_db()
        self.assertTrue(outro.is_active)
        senha = reverse("admin:auth_user_password_change", args=[outro.pk])
        self.assertEqual(self.client.get(senha).status_code, 403)

        imp = usuario("Impresso", "Impresso")  # grupo fora dos permitidos: grupos dele ficam readonly
        self.assertNotIn("groups", self.form_de(imp).fields)
        self.assertEqual(self.post_usuario(imp, first_name="Ana").status_code, 302)
        self.assertEqual(list(imp.groups.values_list("name", flat=True)), ["Impresso"])


class BloqueioMatrizTest(Base):
    """Com a MEC travada por outro usuário, nem a MEC nem o descritor (com tarefas) salvam."""

    def test_post_recusado_com_trava_alheia(self):
        from nucleo.servicos import bloqueio

        bia = usuario("Pedagógica", "Pedagógica")
        self.assertTrue(bloqueio.adquirir(self.mec, bia, "aba-bia"))
        d = Descritor.objects.get(matriz_etapa_cc=self.mec, codigo="D01")
        for url, campo, valor in (
            (reverse("admin:nucleo_matrizetapacc_change", args=[self.mec.pk]), "passo", "77"),
            (reverse("admin:nucleo_descritor_change", args=[d.pk]), "tarefa_set-0-texto", "intrusa"),
        ):
            resp = self.client.get(url)
            self.assertContains(resp, "Somente leitura")
            dados = {**dados_do_form(resp), campo: valor, "_bloqueio_sessao": "aba-root", "confirmar_impacto": "on"}
            if campo.startswith("tarefa_set"):
                dados.update({"tarefa_set-TOTAL_FORMS": "1", "tarefa_set-INITIAL_FORMS": "1",
                              "tarefa_set-0-id": d.tarefa_set.get().pk, "tarefa_set-0-descritor": d.pk,
                              "tarefa_set-0-inicio": "100", "tarefa_set-0-fim": "125", "tarefa_set-0-ordem": "1"})
            self.client.post(url, dados)
        self.mec.refresh_from_db()
        self.assertEqual(self.mec.passo, 25)
        self.assertEqual(d.tarefa_set.get().texto, "t")


class ValidacaoTest(Base):
    def test_nome_de_matriz_repetido_na_edicao(self):
        outra = Matriz.objects.create(ano=self.a26, nome="Exceção")
        url = reverse("admin:nucleo_matriz_change", args=[outra.pk])
        resp = self.client.post(url, {"nome": "Comunitária", "matrizetapacc_set-TOTAL_FORMS": "0",
                                      "matrizetapacc_set-INITIAL_FORMS": "0"})  # fmt: skip
        self.assertContains(resp, "Já existe uma matriz com este nome neste ano.")

    def test_codigo_de_descritor_repetido(self):
        d = Descritor.objects.get(matriz_etapa_cc=self.mec, codigo="D02")
        url = reverse("admin:nucleo_descritor_change", args=[d.pk])
        dados = {**dados_do_form(self.client.get(url)), "codigo": "D01"}
        self.assertContains(self.client.post(url, dados), "Já existe um descritor com este código")

    def test_chaves_readonly_na_edicao(self):
        e5 = self.e["5EF"]
        form = self.client.get(reverse("admin:nucleo_etapa_change", args=[e5.pk])).context["adminform"].form
        self.assertEqual(set(form.fields), {"ordem"})  # tem MEC: escala_max também fica readonly
        form = self.client.get(reverse("admin:nucleo_etapa_change", args=[self.e["3EM"].pk])).context["adminform"].form
        self.assertEqual(set(form.fields), {"escala_max", "ordem"})
        cor = CorPadrao.objects.get(ano=self.a26, padrao=2)
        form = self.client.get(reverse("admin:nucleo_corpadrao_change", args=[cor.pk])).context["adminform"].form
        self.assertEqual(set(form.fields), {"cor", "nome", "confirmar_impacto"})  # T07
        comp = self.c["LP"]
        form = self.client.get(reverse("admin:nucleo_componente_change", args=[comp.pk])).context["adminform"].form
        self.assertEqual(set(form.fields), {"nome"})

    def test_upload_invalido(self):
        url = reverse("admin:nucleo_matriz_reimportar", args=[self.matriz.pk])
        resp = self.client.post(url, {"arquivo": SimpleUploadedFile("m.pdf", b"%PDF-1.4 lixo")})
        self.assertContains(resp, "não foi possível ler m.pdf")
        resp = self.client.post(url, {"arquivo": SimpleUploadedFile("m.pdf", b"<html>")})
        self.assertContains(resp, "não é um PDF nem uma planilha")
        with mock.patch("nucleo.matriz.LIMITE_UPLOAD", 10):
            resp = self.client.post(url, {"arquivo": SimpleUploadedFile("m.pdf", b"%PDF-" + b"x" * 20)})
        self.assertContains(resp, "Arquivo maior que")

    def test_nomes_longos_nao_estouram_o_motivo_do_historico(self):
        a27 = cadastrar(2027)
        nome = "M" * 200
        Matriz.objects.filter(pk=self.matriz.pk).update(nome=nome)
        url = reverse("admin:nucleo_matriz_duplicar", args=[self.matriz.pk])
        self.assertEqual(self.client.post(url, {"ano": a27.pk, "nome": nome}).status_code, 302)
        nova = Matriz.objects.get(ano=a27)
        self.assertEqual(len(nova.history.get().history_change_reason), 100)
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = planilha(Path(tmp) / "m.xlsx", [("MT", "5EF", [(0, 200), (200, 500)], [
                ["D01", "Contar.", "", "–", "Contar", "", "", "", ""],
            ])]).read_bytes()  # fmt: skip
        url = reverse("admin:nucleo_matriz_reimportar", args=[nova.pk])
        resp = self.client.post(url, {"arquivo": SimpleUploadedFile("a" * 150 + ".xlsx", xlsx)})
        self.assertEqual(resp.status_code, 302)


class GradeConcorrenteTest(TransactionTestCase):
    """Duas grades gravam o 1º corte do mesmo parceiro ao mesmo tempo: uma grava, a outra
    recebe conflito (sem IntegrityError do ParceiroAno nem 500)."""

    def test_primeiro_corte_simultaneo(self):
        ano = cadastrar(2026)
        e5 = Etapa.objects.get(ano=ano, codigo="5EF")
        lp = Componente.objects.get(ano=ano, codigo="LP")
        go = Parceiro.objects.create(nome="GO")
        root = usuario("root", is_superuser=True)
        url = reverse("admin:nucleo_corte_changelist") + "?ano=2026"
        clientes = [Client(), Client()]
        envios = []
        for i, cliente in enumerate(clientes):
            cliente.force_login(root)
            dados = dados_da_grade(cliente.get(url))
            dados[f"c-{go.pk}-{e5.pk}-{lp.pk}"] = f"{100 + i}"
            envios.append(dados)
        barreira, status = threading.Barrier(2), []

        def enviar(i):
            try:
                barreira.wait()
                status.append(clientes[i].post(url, envios[i]).status_code)
            finally:
                connection.close()

        threads = [threading.Thread(target=enviar, args=(i,)) for i in range(2)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        self.assertEqual(sorted(status), [200, 302])
        self.assertEqual(ParceiroAno.objects.count(), 1)
        self.assertEqual(Corte.objects.count(), 1)
