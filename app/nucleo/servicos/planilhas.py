"""Planilhas do impresso (S13): F1 por etapa e F2, num zip ("Baixar planilhas").

F1: porte de `api/src/modules/planilhas/{layout,service}.ts` para openpyxl, generalizado para
N padrões; lê o `f1.dados[etapa]` do resolvido (o mesmo JSON do JS, D-130), então os rótulos dos
padrões já vêm resolvidos (`pdNN` do projeto ou nome do ano, D-175).
F2: uma linha por seção (seção, título, texto), com o texto puro.
"""

import io
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from nucleo import texto
from nucleo.models import MatrizEtapaCC, ProjetoPasso
from nucleo.servicos.empacotar import nome_zip
from nucleo.servicos.resolver import VAZIO, resolver

# Colunas fixas por componente, copiadas dos modelos do núcleo de impressão (layout.ts).
LAYOUTS = {
    "LP": {
        "inicio": ["Etapa", "Descritor", "Descrição da habilidade", "Pre requisito"],
        "fim": ["Praticas de linguagem:", "Objeto de conhecimento:", "Habilidades relacionadas", "Habilidade correspondente"],
        "vazio": "–",
        "larguras": ([8, 12, 50, 30], 28, [30, 30, 30, 40]),
    },
    "MT": {
        "inicio": ["Etapa", "Cód. Descritor", "Descritor", "Pre requisito"],
        "fim": ["Eixo cognitivo (Saeb):", "Unidade temática:", "Habilidades relacionadas:", "Habilidade correspondente:"],
        "vazio": "...",
        "larguras": ([8, 14, 50, 30], 28, [30, 30, 30, 40]),
    },
}
AMARELO = PatternFill("solid", fgColor="FFFFFF00")  # COR_NIVEL
NEGRITO = Font(bold=True)
TOPO = Alignment(vertical="top", wrap_text=True)


def descritores(dados):
    """[(código, descritor)] de uma etapa do resolvido, na ordem da matriz."""
    return [(c, d) for c, d in dados.items() if isinstance(d, dict) and "scale" in d]


def niveis(dados):
    """Níveis distintos de todas as linhas ∪ mínimo, crescentes; `fim` = próximo nível ou máximo."""
    faixa = dados["scaleRange"]
    todos = {faixa["min"]} | {linha["level"] for _, d in descritores(dados) for ls in d["scale"].values() for linha in ls}
    ordenados = sorted(todos)
    return [(n, ordenados[i + 1] if i + 1 < len(ordenados) else faixa["max"]) for i, n in enumerate(ordenados)]


def celula_do_nivel(linhas, vazio):
    """Conteúdo das linhas sem os "---", unidos por "#"; nada → `vazio`."""
    return "#".join(c for c in (linha["content"] for linha in linhas) if c != VAZIO) or vazio


def separar_habilidades(skills):
    """Itens entre parênteses no início são a habilidade correspondente; o resto, as relacionadas."""
    return "\n".join(s for s in skills if not s.startswith("(")), "\n".join(s for s in skills if s.startswith("("))


def nome_planilha(nome_projeto, componente, etapa):
    return f"{re.sub(r'[/\\:*?\"<>|]', '', nome_projeto).strip()} - {componente} {etapa.upper()}.xlsx"


def _bytes(wb):
    for ws in wb:
        for linha in ws.iter_rows():
            for c in linha:
                if isinstance(c.value, str):
                    c.data_type = "s"  # "=HYPERLINK(...)" num texto não vira fórmula (como admin/cortes.py)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _larguras(ws, larguras):
    for i, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = largura


def planilha_f1(dados, etapa, componente, cores, passo=None):
    """xlsx de uma etapa: abas "Níveis" (descritor × nível) e "Escala" (padrões, limites e cor)."""
    layout = LAYOUTS.get(componente, LAYOUTS["LP"])
    ns = niveis(dados)
    padroes = [p["chave"] for p in dados["padroes"]]
    wb = Workbook()
    ws = wb.active
    ws.title = "Níveis"
    ws.freeze_panes = "A2"
    ws.append([*layout["inicio"], *(f"{a} a {b}" for a, b in ns), *layout["fim"]])
    inicio, nivel, fim = layout["larguras"]
    _larguras(ws, [*inicio, *[nivel] * len(ns), *fim])
    for c in ws[1]:
        c.font = NEGRITO
    for codigo, d in descritores(dados):
        relacionadas, correspondente = separar_habilidades(d["bncc"]["skills"])
        linhas = [linha for p in padroes for linha in d["scale"].get(p, [])]
        ws.append([
            etapa.upper(), codigo, d["description"], "#".join(d["prerequisites"]),
            *(celula_do_nivel([x for x in linhas if x["level"] == n], layout["vazio"]) for n, _ in ns),
            d["bncc"]["practices"], d["bncc"]["knowledge"], relacionadas, correspondente,
        ])  # fmt: skip
        for c in ws[ws.max_row]:
            c.alignment = TOPO
    for linha in ws.iter_rows(min_col=5, max_col=4 + len(ns)):
        for c in linha:
            c.fill = AMARELO

    ws = wb.create_sheet("Escala")
    faixa = dados["scaleRange"]
    for linha in (["Etapa", etapa.upper()], ["Mínimo", faixa["min"]], ["Máximo", faixa["max"]], ["Intervalo mínimo", "" if passo is None else passo], []):
        ws.append(linha)
    ws.append(["Padrão", "Início (≥)", f"Fim (< / ≤ no {len(padroes):02d})", "Cor"])
    for c in ws[ws.max_row]:
        c.font = NEGRITO
    limites = [faixa["min"], *(dados["cortes"].get(f"padrao-{i}") for i in range(1, len(padroes))), faixa["max"]]
    for i, p in enumerate(dados["padroes"]):
        cor = cores.get(p["chave"], "")
        ws.append([p["rotulo"], limites[i], limites[i + 1], cor])
        if cor:
            ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor="FF" + cor[1:].upper())
    _larguras(ws, [20, 14, 20, 12])
    for (c,) in ws.iter_rows(max_col=1):
        if c.value is not None:
            c.font = NEGRITO
    return _bytes(wb)


def planilha_f2(secoes):
    """xlsx do F2: uma linha por seção, com título e texto puros."""
    wb = Workbook()
    ws = wb.active
    ws.title = "F2"
    ws.freeze_panes = "A2"
    ws.append(["Seção", "Título", "Texto"])
    for c in ws[1]:
        c.font = NEGRITO
    for s in secoes or []:
        ws.append([s["id"], texto.html_para_texto(s["titulo"]), texto.html_para_texto(s["corpo"])])
        for c in ws[ws.max_row]:
            c.alignment = TOPO
    _larguras(ws, [16, 40, 100])
    return _bytes(wb)


def passos(revista):
    """{etapa: passo} como no resolver (D-053, D-091): a linha de ProjetoPasso sobrescreve a MEC."""
    pa, filtro = revista.projeto.parceiro_ano, {"componente_id": revista.componente_id}
    saida = dict(MatrizEtapaCC.objects.filter(matriz_id=pa.matriz_id, **filtro).values_list("etapa__codigo", "passo"))
    saida.update(ProjetoPasso.objects.filter(projeto=revista.projeto, **filtro).values_list("etapa__codigo", "passo"))
    return {etapa.lower(): passo for etapa, passo in saida.items()}


def zip_planilhas(revista):
    """(nome, bytes) do zip com as planilhas F1 (uma por etapa) e F2. Revista de `resolver.carregar`."""
    f1, f2 = (resolvido := resolver(revista))["f1"], resolvido["f2"]
    nome, cc, por_etapa = revista.projeto.nome, revista.componente.codigo, passos(revista)
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as z:
        for etapa in f1["etapas"]:
            z.writestr(nome_planilha(nome, cc, etapa), planilha_f1(f1["dados"][etapa], etapa, cc, f1["cores"], por_etapa.get(etapa)))
        z.writestr(nome_planilha(nome, cc, "F2"), planilha_f2(f2["secoes"]))
    return f"planilhas-{nome_zip(revista)}", saida.getvalue()
