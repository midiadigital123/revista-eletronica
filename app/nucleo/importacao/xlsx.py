"""Lê a matriz de uma planilha no layout da exportação (`api/src/modules/planilhas`).

Cada aba "Níveis" (A1 "Etapa", B1 com o cabeçalho do código) é uma etapa × CC. Colunas: Etapa, código,
descrição, pré-requisitos (separados por "#"), uma coluna por intervalo "125 a 150"
(tarefas separadas por "#"; "–", "..." ou "---" = sem tarefa) e as 4 colunas da BNCC.
O CC sai do cabeçalho da 2ª coluna ("Descritor" na LP, "Cód. Descritor" na MT).
A aba "Escala" (cortes e cores do projeto) é ignorada: cortes não são da matriz (D-059).

Devolve o mesmo formato de `importacao.pdf.extrair`.
"""
import re
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

CC_DO_CABECALHO = {"Descritor": "LP", "Cód. Descritor": "MT"}
INTERVALO = re.compile(r"^\s*(\d+)\s+a\s+(\d+)\s*$")
VAZIOS = {"", "–", "...", "---"}


def _texto(v):
    return "" if v is None else str(v).strip()


def _aba(ws):
    linhas = ws.iter_rows(values_only=True)
    cab = [_texto(c) for c in next(linhas, ())]
    cc = CC_DO_CABECALHO[cab[1]]
    intervalos = [(i, *map(int, m.groups())) for i, c in enumerate(cab) if (m := INTERVALO.match(c))]
    if not intervalos:
        raise ValueError(f'aba "{ws.title}": nenhuma coluna de intervalo ("125 a 150")')
    for k, (i, inicio, fim) in enumerate(intervalos):
        if inicio >= fim:
            raise ValueError(f'aba "{ws.title}": intervalo "{cab[i]}" com início ≥ fim')
        if k and (i != intervalos[k - 1][0] + 1 or inicio != intervalos[k - 1][2]):
            raise ValueError(f'aba "{ws.title}": "{cab[i]}" não continua "{cab[intervalos[k - 1][0]]}"')
    fim_bncc = intervalos[-1][0] + 1  # praticas, conhecimento, relacionadas, correspondente
    if len(cab) < fim_bncc + 4:
        raise ValueError(f'aba "{ws.title}": faltam colunas da BNCC depois dos intervalos')
    etapas, descritores, codigos = set(), [], set()
    for n, linha in enumerate(linhas, start=2):
        v = [_texto(c) for c in linha]
        v += [""] * (fim_bncc + 4 - len(v))
        if not any(v):
            continue
        if not v[1]:
            raise ValueError(f'aba "{ws.title}" linha {n}: descritor sem código')
        if v[1] in codigos:
            raise ValueError(f'aba "{ws.title}" linha {n}: código {v[1]} repetido')
        codigos.add(v[1])
        etapas.add(v[0].upper())
        bncc = [" ".join(s.split()) for s in v[fim_bncc : fim_bncc + 4]]
        descritores.append({
            "codigo": v[1],
            "descricao": v[2],
            "prerequisitos": [p.strip() for p in v[3].split("#") if p.strip()],
            "praticas": bncc[0],
            "conhecimento": bncc[1],
            "habilidades_relacionadas": bncc[2],
            "habilidade_correspondente": bncc[3],
            # De cima para baixo, como no PDF.
            "tarefas": [
                {"inicio": inicio, "fim": fim, "texto": t.strip()}
                for i, inicio, fim in reversed(intervalos)
                for t in v[i].split("#")
                if t.strip() not in VAZIOS
            ],
        })
    if len(etapas) != 1:
        raise ValueError(f'aba "{ws.title}": esperada uma etapa, achadas {sorted(etapas) or "nenhuma"}')
    escala = {"min": intervalos[0][1], "max": intervalos[-1][2]}
    return (cc, etapas.pop()), {"escala": escala, "descritores": descritores}


def extrair(arquivo):
    try:
        wb = load_workbook(arquivo, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException) as e:
        raise ValueError("arquivo não é um .xlsx válido") from e
    saida = {}
    try:
        for ws in wb.worksheets:
            cab = [_texto(c) for c in next(ws.iter_rows(max_row=1, max_col=2, values_only=True), ())]
            if len(cab) < 2 or cab[0] != "Etapa" or cab[1] not in CC_DO_CABECALHO:
                continue  # aba "Escala" (A1 também é "Etapa", B1 é o valor) e outras
            chave, dados = _aba(ws)
            if chave in saida:
                raise ValueError(f"{chave[0]} {chave[1]} aparece em mais de uma aba")
            saida[chave] = dados
    finally:
        wb.close()
    if not saida:
        raise ValueError('nenhuma aba no layout da exportação (A1 "Etapa", B1 "Descritor" ou "Cód. Descritor")')
    return saida
