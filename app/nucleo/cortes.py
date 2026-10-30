"""Grade de cortes por ano e planilha de grupos (S11a; D-049, D-055, D-056, D-135, D-168).

Regra comum às duas telas da grade: a do admin (`admin/cortes.py`) e a página própria
(`urls/cortes.py`, protótipo da opção B). A grade é parceiro × (etapa × CC) de um ano; cada
célula aceita "350, 400, 500". Só as linhas (parceiros) que mudaram são gravadas, cada uma
inteira por `gravar_cortes` (célula vazia = corte apagado), tudo numa transação: um erro, nada grava.
"""

import io
import re

from django.contrib import messages
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from openpyxl import Workbook

from nucleo.models import Componente, Corte, Etapa, Parceiro, ParceiroAno
from nucleo.servicos.cascata import gravar_cortes, validar_linhas

SEPARADOR = re.compile(r"[\s,;]+")
SEP = "|"
CONFLITO = (
    "Os cortes deste parceiro mudaram desde que a página foi aberta (agora: {}). "
    "Seus valores foram mantidos; confira e salve de novo para sobrescrever."
)


def formatar(valores):
    return ", ".join(map(str, valores or ()))


def interpretar(texto):
    """"350, 400 500" -> [350, 400, 500]; vazio -> []. Não inteiro: ValueError."""
    partes = [p for p in SEPARADOR.split(texto.strip()) if p]
    if any(not (p.isascii() and p.lstrip("-").isdigit()) for p in partes):
        raise ValueError
    return [int(p) for p in partes]


def formatar_texto(texto):
    """Normaliza o texto da célula ("350 400" -> "350, 400") para comparar com o original."""
    try:
        return formatar(interpretar(texto))
    except ValueError:
        return texto


def numerar_grupos(listas):
    """D-056: listas distintas, ordenadas corte a corte, numeradas a partir de 1."""
    return {v: i for i, v in enumerate(sorted(set(map(tuple, listas))), start=1)}


def planilha(ano):
    """xlsx com uma aba "ETAPA - CC" por etapa × CC do ano: Grupo, Parceiro, Etapa, CC.
    O nº do grupo aparece só na 1ª linha do grupo, como na planilha do impresso."""
    abas = {
        (e.ordem, e.codigo, c.codigo): []
        for e in Etapa.objects.filter(ano=ano)
        for c in Componente.objects.filter(ano=ano)
    }
    for c in Corte.objects.filter(ano=ano).select_related("parceiro_ano__parceiro", "etapa", "componente"):
        abas[(c.etapa.ordem, c.etapa.codigo, c.componente.codigo)].append((c.parceiro_ano.parceiro.nome, c.valores))
    wb = Workbook()
    wb.remove(wb.active)
    for (_, etapa, cc), itens in sorted(abas.items()):
        ws = wb.create_sheet(re.sub(r"[/\\?*\[\]:]", "-", f"{etapa} - {cc}")[:31])
        ws.append(["Grupo", "Parceiro", "Etapa", cc])
        grupos = numerar_grupos(v for _, v in itens)
        anterior = None
        for nome, valores in sorted(itens, key=lambda i: (grupos[tuple(i[1])], i[0])):
            g = grupos[tuple(valores)]
            ws.append([g if g != anterior else None, nome, etapa, formatar(valores)])
            anterior = g
        for linha in ws.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str):
                    celula.data_type = "s"  # "=HYPERLINK(...)" no nome do parceiro não vira fórmula
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def pode_ver(user):
    return user.has_perm("nucleo.view_corte") or user.has_perm("nucleo.change_corte")


def pode_gravar(user):
    return user.has_perms(["nucleo.add_corte", "nucleo.change_corte"])


class Conflito(Exception):
    def __init__(self, parceiro, agora):
        super().__init__(f"os cortes de {parceiro} mudaram desde que a página foi aberta.")
        self.parceiro, self.agora = parceiro, agora


def linhas_atuais(ano, colunas, parceiro_ids=None):
    """{parceiro_id: [texto de cada coluna]} com os cortes gravados no ano."""
    indice = {(e.pk, c.pk): i for i, (e, c) in enumerate(colunas)}
    cortes = Corte.objects.filter(ano=ano).select_related("parceiro_ano")
    if parceiro_ids is not None:
        cortes = cortes.filter(parceiro_ano__parceiro_id__in=parceiro_ids)
    linhas = {}
    for c in cortes:
        linha = linhas.setdefault(c.parceiro_ano.parceiro_id, [""] * len(colunas))
        linha[indice[(c.etapa_id, c.componente_id)]] = formatar(c.valores)
    return linhas


def grade(request, ano):
    """Monta a grade do ano e, no POST, grava. Devolve (contexto, gravou).

    `gravou` verdadeiro: o chamador redireciona (a mensagem de sucesso já foi posta).
    POST sem `add_corte` + `change_corte`: PermissionDenied.
    """
    etapas = list(Etapa.objects.filter(ano=ano).order_by("ordem"))
    componentes = list(Componente.objects.filter(ano=ano).order_by("codigo"))
    colunas = [(e, c) for e in etapas for c in componentes]
    parceiros = list(Parceiro.objects.order_by("nome"))
    atuais = linhas_atuais(ano, colunas)
    pode = pode_gravar(request.user)
    post = request.POST if request.method == "POST" else None
    if post is not None and not pode:
        raise PermissionDenied

    linhas, mudadas, com_erro = [], {}, False
    for p in parceiros:
        atual = atuais.get(p.pk, [""] * len(colunas))
        # Um "o-" por parceiro (a linha inteira como estava na página): 1 campo a menos por
        # célula, para caber no DATA_UPLOAD_MAX_NUMBER_FIELDS (1000) com ~100 parceiros.
        original = None if post is None else post.get(f"o-{p.pk}")
        erros, celulas, valores = [], [], []
        for (e, c), texto_atual in zip(colunas, atual):
            chave = f"{p.pk}-{e.pk}-{c.pk}"
            texto = texto_atual if original is None else post.get(f"c-{chave}")
            if texto is None:  # célula faltando no POST não é "apagar"
                erros.append(f"{e.codigo} {c.codigo}: célula ausente no envio. Recarregue a página.")
                texto = ""
            texto = texto.strip()
            celulas.append({"chave": chave, "texto": texto, "titulo": f"{p} {e.codigo} {c.codigo}", "inicio": c == componentes[0]})
            try:
                lista = interpretar(texto)
            except ValueError:
                erros.append(f"{e.codigo} {c.codigo}: cortes precisam ser inteiros (ex.: 350, 400, 500).")
                continue
            if lista:
                valores.append((e, c, lista))
        # Parceiro criado depois de a página abrir não tem "o-": fica como está.
        nova = [cel["texto"] for cel in celulas]
        if original is not None and (erros or SEP.join(map(formatar_texto, nova)) != original):
            if not erros and original != SEP.join(atual):
                erros.append(CONFLITO.format(SEP.join(atual).replace(SEP, " | ") or "vazia"))
            elif not erros:
                try:
                    validar_linhas(valores, ano.pk)
                except ValidationError as ex:
                    erros.extend(ex.messages)
            if not erros:
                mudadas[p.pk] = (p, valores, original)
        com_erro |= bool(erros)
        for cel, (e, c) in zip(celulas, colunas):  # mensagens começam por "ETAPA CC:" (cascata.validar_linhas)
            cel["invalida"] = any(m.startswith(f"{e.codigo} {c.codigo}:") for m in erros)
        # Com conflito, o "o-" passa a ser o estado atual: salvar de novo sobrescreve de propósito.
        linhas.append({"parceiro": p, "celulas": celulas, "erros": erros, "original": SEP.join(atual)})

    contexto = {"etapas": etapas, "componentes": componentes, "colunas": colunas, "linhas": linhas, "pode_gravar": pode}
    if post is not None and not com_erro:
        try:
            gravar(request.user, ano, colunas, mudadas)
        except Conflito as ex:
            messages.error(request, f"Nada foi gravado: {ex}")
            for linha in linhas:
                if linha["parceiro"].pk == ex.parceiro.pk:
                    linha["erros"].append(CONFLITO.format(ex.agora.replace(SEP, " | ") or "vazia"))
                    linha["original"] = ex.agora
        else:
            messages.success(request, f"Cortes gravados ({len(mudadas)} parceiro(s) alterado(s)).")
            return contexto, True
    elif com_erro:
        messages.error(request, "Nada foi gravado: corrija os erros marcados.")
    return contexto, False


def gravar(user, ano, colunas, mudadas):
    with transaction.atomic():
        pas = {}
        for pk in sorted(mudadas):
            # get_or_create: 1º corte simultâneo do mesmo parceiro não vira IntegrityError.
            pas[pk] = ParceiroAno.objects.get_or_create(parceiro_id=pk, ano=ano)[0]
        # Trava em ordem de pk (sem deadlock entre duas grades) e relê sob a trava.
        list(ParceiroAno.objects.select_for_update(no_key=True).filter(pk__in=[pa.pk for pa in pas.values()]).order_by("pk"))
        agora = linhas_atuais(ano, colunas, parceiro_ids=list(mudadas))
        for pk, (p, valores, original) in sorted(mudadas.items()):
            atual = SEP.join(agora.get(pk, [""] * len(colunas)))
            if atual != original:
                raise Conflito(p, atual)
            gravar_cortes(pas[pk], valores, user)
